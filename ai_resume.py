import streamlit as st
import pdfplumber
import docx
import chromadb
from fpdf import FPDF
from langchain_groq import ChatGroq
from langchain_community.vectorstores import Chroma
from langchain_huggingface import HuggingFaceEmbeddings

# Initialize LLM
llm = ChatGroq(
    groq_api_key="",
    model="llama3-8b-8192",
    temperature=0,
    max_tokens=None,
    timeout=None,
    max_retries=2,
)

# Initialize ChromaDB
chroma_client = chromadb.PersistentClient(path="./chroma_db")
embedding_function = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")
vectorstore = Chroma(client=chroma_client, embedding_function=embedding_function)

# Extract text from PDF
def extract_text_from_pdf(pdf_file):
    with pdfplumber.open(pdf_file) as pdf:
        text = "\n".join(page.extract_text(layout=True) for page in pdf.pages if page.extract_text(layout=True))
    return text

# Extract text from DOCX
def extract_text_from_docx(docx_file):
    doc = docx.Document(docx_file)
    text = "\n".join(para.text.strip() for para in doc.paragraphs if para.text.strip())
    return text

# Generate recommended projects
def recommend_projects(job_description, existing_projects):
    prompt = f"""
    Recommend 2-3 projects relevant to the following job description:
    {job_description}
    Provide a brief explanation of each project.
    The output should be formatted as a list of projects, each having a title and a short description.
    """
    recommended_projects = llm.invoke(prompt).content
    
    # Combine user-provided and recommended projects
    final_projects = existing_projects + "\n\n" + recommended_projects if existing_projects else recommended_projects
    return final_projects

# Remove empty sections from the resume
def remove_empty_sections(resume_text):
    sections = resume_text.split("--------------------------------------------------------")
    cleaned_sections = []
    
    for section in sections:
        if "[" not in section and section.strip():  # Remove sections with placeholders or empty content
            cleaned_sections.append(section.strip())
    
    return "\n--------------------------------------------------------\n".join(cleaned_sections)

# Generate a new ATS-friendly resume with projects
def generate_resume(details, job_description, projects):
    final_projects = recommend_projects(job_description, projects)
    
    prompt = f"""
    Generate a professional ATS-friendly resume based on the following details:
    {details}
    Ensure the 'Projects' section includes the following:
    {final_projects}
    Tailor the resume to align with the following job description:
    {job_description}
    Ensure the output follows a structured professional resume format with sections for Contact Information, Summary, Education, Skills, Experience, Projects, Achievements, and Certifications.
    If any details are missing, skip that section.
    Strictly remove "Here's an ATS-friendly resume based on the provided details:" line at the beginning and any note section placed at the end.
    Make the headings bold and the content in normal font. Draw a line after each section.
    """
    resume_text = llm.invoke(prompt).content
    return remove_empty_sections(resume_text)

# Modify existing resume while keeping original format
def modify_resume(text, job_description, projects):
    final_projects = recommend_projects(job_description, projects)
    
    prompt = f"""
    Modify the following resume to be ATS-friendly according to the given job description:
    Resume Text:
    {text}
    Ensure the 'Projects' section includes the following:
    {final_projects}
    Job Description:
    {job_description}
    Ensure the original format, structure, and layout remain unchanged. Only update the content to be more ATS-friendly by improving keyword relevance and alignment with the job description.
    Do not add additional content or sections not present in the original resume.
    Strictly remove "Here's an ATS-friendly resume based on the provided details:" line at the beginning and any note section placed at the end.
    Make the headings bold and visible. Draw a line after each section.
    """
    modified_text = llm.invoke(prompt).content
    return remove_empty_sections(modified_text)

# Convert text to PDF
def text_to_pdf(text, filename):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=12)
    
    text = text.encode("latin-1", "ignore").decode("latin-1")
    pdf.multi_cell(0, 10, text)
    pdf.output(filename, "F")

# Streamlit UI
st.title("AI Resume Generator")

option = st.radio("Choose an option:", ["Upload Basic Details & Job Description", "Upload Resume & Job Description"])

if option == "Upload Basic Details & Job Description":
    st.subheader("Enter Your Details")
    name = st.text_input("Full Name")
    email = st.text_input("Email")
    phone = st.text_input("Phone Number")
    education = st.text_area("Education (School, College, Degree)")
    skills = st.text_area("Skills (comma-separated)")
    experience = st.text_area("Experience")
    projects = st.text_area("Projects (with a small description)")
    achievements = st.text_area("Achievements")
    certifications = st.text_area("Certifications")
    job_description = st.text_area("Job Description")

    if st.button("Generate Resume"):
        user_details = {
            "name": name,
            "email": email,
            "phone": phone,
            "education": education,
            "skills": skills.split(","),
            "experience": experience,
            "projects": projects,
            "achievements": achievements,
            "certifications": certifications,
        }
        resume_text = generate_resume(user_details, job_description, projects)
        pdf_filename = "generated_resume.pdf"
        text_to_pdf(resume_text, pdf_filename)
        st.subheader("Generated Resume")
        st.download_button("Download Resume", open(pdf_filename, "rb"), file_name=pdf_filename, mime="application/pdf")
